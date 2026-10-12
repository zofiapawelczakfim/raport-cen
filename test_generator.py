"""Test end-to-end: 3 dni publikacji, zmiana ceny, sprzedaż lokalu, przełom roku, kontrola XML/MD5/CSV.
Uruchom:  python testy/test_generator.py
"""
import csv
import hashlib
import shutil
import sys
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "generator"))
import generuj as g  # noqa: E402

BASE = "https://przyklad.github.io/raport-cen"
NSX = "{urn:otwarte-dane:harvester:1.13}"
KOLEJNOSC_DATASET = ["extIdent", "title", "description", "url", "updateFrequency", "hasDynamicData",
                     "hasHighValueData", "hasHighValueDataFromEuropeanCommissionList", "hasResearchData",
                     "categories", "resources", "tags"]
KOLEJNOSC_RESOURCE = ["extIdent", "url", "title", "description", "availability", "dataDate", "specialSigns",
                      "hasDynamicData", "hasHighValueData", "hasHighValueDataFromEuropeanCommissionList",
                      "hasResearchData", "containsProtectedData"]


def uruchom(tmp, xlsx, dzien, oczekiwany=0):
    kod = g.main(["--zrodlo", str(xlsx), "--data", dzien, "--wyjscie", str(tmp / "docs"),
                  "--stan", str(tmp / "stan"), "--base-url", BASE])
    assert kod == oczekiwany, f"{dzien}: kod {kod}, oczekiwano {oczekiwany}"


def czytaj_csv(tmp, dzien):
    p = tmp / "docs" / "csv" / dzien[:4] / g.nazwa_csv("przyklad-deweloper", dzien)
    with open(p, encoding="utf-8", newline="") as f:
        return list(csv.reader(f))


def ustaw(xlsx, wiersz, kol, wartosc):
    wb = openpyxl.load_workbook(xlsx)
    wb["Oferta"].cell(row=wiersz, column=kol, value=wartosc)
    wb.save(xlsx)


def main():
    tmp = Path(tempfile.mkdtemp())
    xlsx = tmp / "cennik.xlsx"
    shutil.copy(ROOT / "szablon" / "Cennik_szablon.xlsx", xlsx)
    # przy migracji: lokal A1 ma cenę obowiązującą od 2026-09-01
    ustaw(xlsx, 2, 17, "2026-09-01")

    # Dzień 1
    uruchom(tmp, xlsx, "2026-12-30")
    r = czytaj_csv(tmp, "2026-12-30")
    assert r[0] == g.NAGLOWKI and len(r) == 4
    assert all(len(w) == 58 for w in r)
    assert all(c.strip() != "" for w in r[1:] for c in w), "puste komórki – powinno być X"
    a1, a2, mp7 = r[1], r[2], r[3]
    assert a1[36] == "A1" and a1[37] == "12500" and a1[38] == "2026-09-01", a1[36:39]
    assert a2[38] == "2026-12-30"
    assert a1[2] == "0000123456" and a1[17] == "61-888" and a1[57].startswith("https://")
    assert mp7[35] == "X" and mp7[44] == "MP7" and mp7[45] == "39000" and mp7[46] == "2026-12-30"

    # Dzień 2: podwyżka A2, A1 bez zmian
    ustaw(xlsx, 3, 4, 13100)
    ustaw(xlsx, 3, 5, 825300)
    uruchom(tmp, xlsx, "2026-12-31")
    r = czytaj_csv(tmp, "2026-12-31")
    assert r[1][38] == "2026-09-01", "data A1 nie powinna się zmienić"
    assert r[2][37] == "13100" and r[2][38] == "2026-12-31" and r[2][40] == "2026-12-31"

    # Dzień 3 (nowy rok): A1 sprzedany -> usunięty wiersz
    wb = openpyxl.load_workbook(xlsx)
    wb["Oferta"].delete_rows(2)
    wb.save(xlsx)
    uruchom(tmp, xlsx, "2027-01-01")
    r = czytaj_csv(tmp, "2027-01-01")
    assert len(r) == 3 and r[1][36] == "A2" and r[1][38] == "2026-12-31"

    # Ponowne uruchomienie tego samego dnia – nic się nie dubluje
    uruchom(tmp, xlsx, "2027-01-01")

    # XML
    xml_p = tmp / "docs" / "przyklad-deweloper.xml"
    xml = xml_p.read_bytes()
    md5 = (tmp / "docs" / "przyklad-deweloper.md5").read_text()
    assert md5 == hashlib.md5(xml).hexdigest() and md5 == md5.lower() and len(md5) == 32
    root = ET.fromstring(xml)
    assert root.tag == NSX + "datasets"
    ds = list(root)
    assert [d.find("extIdent").text for d in ds] == ["przyklad-deweloper-ceny-2026", "przyklad-deweloper-ceny-2027"]
    for d in ds:
        assert d.get("status") == "published"
        assert [c.tag for c in d] == KOLEJNOSC_DATASET, [c.tag for c in d]
        assert len(d.find("extIdent").text) <= 36
        for t in ("title", "description"):
            assert [c.tag for c in d.find(t)] == ["polish", "english"]
        for res in d.find("resources"):
            assert res.get("status") == "published"
            assert [c.tag for c in res] == KOLEJNOSC_RESOURCE, [c.tag for c in res]
            assert len(res.find("extIdent").text) <= 36
            assert res.find("availability").text == "local"
            assert res.find("specialSigns/specialSign").text == "X"
            assert " " not in res.find("url").text
    ids = [res.find("dataDate").text for d in ds for res in d.find("resources")]
    assert ids == ["2026-12-30", "2026-12-31", "2027-01-01"], ids
    for d in ds:
        for el in d.iter():
            if el.tag in ("polish", "english"):
                lim = 300 if el in list(d.find("title")) else 1000
                assert len(el.text) <= lim

    # Walidacja błędów: nieznana inwestycja
    ustaw(xlsx, 2, 1, "Nie ma takiej")
    kod = g.main(["--zrodlo", str(xlsx), "--tylko-sprawdz"])
    assert kod == 2

    # Ochrona identyfikatora: zmiana ID = błąd
    ustaw(xlsx, 2, 1, "Osiedle Zielone")
    wb = openpyxl.load_workbook(xlsx)
    for row in wb["Firma"].iter_rows(min_row=2):
        if row[0].value == g.POLE_ID:
            row[1].value = "inna-nazwa"
    wb.save(xlsx)
    uruchom(tmp, xlsx, "2027-01-02", oczekiwany=2)

    print("WSZYSTKIE TESTY OK")
    print("Przykładowy XML:", xml_p)
    return tmp


if __name__ == "__main__":
    main()
