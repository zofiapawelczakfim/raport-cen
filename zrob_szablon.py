"""Tworzy szablon arkusza cennika (Cennik_szablon.xlsx) zgodny z generatorem."""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "generator"))
import generuj as g  # noqa: E402

F = "Arial"
NAG = PatternFill("solid", fgColor="1F3864")
WYM = PatternFill("solid", fgColor="FFF2CC")  # pola wymagane
OPC = PatternFill("solid", fgColor="F2F2F2")
PRZ = PatternFill("solid", fgColor="E2EFDA")  # przykład
cienka = Side(style="thin", color="BFBFBF")
RAMKA = Border(left=cienka, right=cienka, top=cienka, bottom=cienka)


def naglowek(ws, kolumny, szer):
    for i, (k, s) in enumerate(zip(kolumny, szer), start=1):
        c = ws.cell(row=1, column=i, value=k)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = NAG
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = RAMKA
        ws.column_dimensions[c.column_letter].width = s
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "A2"


def wiersz(ws, r, wartosci, fill=PRZ):
    for i, v in enumerate(wartosci, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=F)
        c.fill = fill
        c.border = RAMKA


def main(cel: Path):
    wb = Workbook()

    # ---------------- Instrukcja
    ws = wb.active
    ws.title = "Instrukcja"
    ws.column_dimensions["A"].width = 110
    linie = [
        ("Cennik do automatycznego raportowania na dane.gov.pl", True),
        ("", False),
        ("Jak to działa: codziennie w nocy automat pobiera ten arkusz i publikuje cennik na dane.gov.pl.", False),
        ("Wystarczy, że ta tabela jest aktualna. Daty zmian cen automat ustawia sam.", False),
        ("", False),
        ("Co uzupełniać:", True),
        ("1. „Firma”: dane dewelopera (raz). Żółte pola są wymagane. Identyfikatora technicznego nie zmieniaj po starcie.", False),
        ("2. „Inwestycje”: jedna linia na inwestycję (adres budowy + link do prospektu informacyjnego).", False),
        ("3. „Oferta”: jedna linia na każdy lokal / dom / osobno sprzedawane miejsce postojowe lub komórkę, które są W SPRZEDAŻY.", False),
        ("   • Sprzedany lub zarezerwowany lokal (wycofany z oferty): usuń wiersz.", False),
        ("   • Zmiana ceny: wpisz nową cenę. Automat sam zapisze datę zmiany i zachowa historię na dane.gov.pl.", False),
        ("   • Miejsce postojowe przypisane do mieszkania: wpisz w tym samym wierszu co mieszkanie.", False),
        ("     Sprzedawane osobno: osobny wiersz bez rodzaju nieruchomości, za to z oznaczeniem części (np. MP7).", False),
        ("   • Puste komórki są OK. Automat wpisze w nie znak „X” (wymóg ministerstwa).", False),
        ("   • Kwoty brutto (z VAT), w zł, bez „zł” i bez spacji. Np. 685000 albo 12450.50.", False),
        ("   • „Ceny obowiązują od”: wypełnij tylko przy przenoszeniu istniejącej oferty (data, od kiedy cena obowiązuje). Później zostaw puste.", False),
        ("", False),
        ("Zielone wiersze to przykład. Usuń je i wpisz swoje dane. Nie zmieniaj nazw zakładek ani nagłówków w 1. wierszu.", True),
    ]
    for r, (t, b) in enumerate(linie, start=1):
        c = ws.cell(row=r, column=1, value=t)
        c.font = Font(name=F, bold=b, size=14 if r == 1 else 11)

    # ---------------- Firma
    ws = wb.create_sheet("Firma")
    naglowek(ws, ["Pole", "Wartość"], [48, 60])
    przyklad = {
        "Nazwa dewelopera": "Przykład Deweloper Sp. z o.o.",
        "Forma prawna": "Spółka z ograniczoną odpowiedzialnością",
        "Nr KRS": "0000123456",
        "Nr NIP": "7791234567",
        "Nr REGON": "301234567",
        "Nr telefonu": "612345678",
        "Adres e-mail": "biuro@przyklad-deweloper.pl",
        "Adres strony internetowej": "https://przyklad-deweloper.pl",
        "Siedziba - województwo": "wielkopolskie",
        "Siedziba - powiat": "Poznań",
        "Siedziba - gmina": "Poznań",
        "Siedziba - miejscowość": "Poznań",
        "Siedziba - ulica": "ul. Półwiejska",
        "Siedziba - nr nieruchomości": "10",
        "Siedziba - nr lokalu": "5",
        "Siedziba - kod pocztowy": "61-888",
        "Biuro sprzedaży - województwo": "wielkopolskie",
        "Biuro sprzedaży - powiat": "Poznań",
        "Biuro sprzedaży - gmina": "Poznań",
        "Biuro sprzedaży - miejscowość": "Poznań",
        "Biuro sprzedaży - ulica": "ul. Półwiejska",
        "Biuro sprzedaży - nr nieruchomości": "10",
        "Biuro sprzedaży - nr lokalu": "5",
        "Biuro sprzedaży - kod pocztowy": "61-888",
        "Sposób kontaktu nabywcy z deweloperem": "telefonicznie, e-mailowo, osobiście w biurze sprzedaży",
        g.POLE_ID: "przyklad-deweloper",
        g.POLE_STRONA_CEN: "https://przyklad-deweloper.pl/ceny",
    }
    pola = [(p, w) for p, _, _, w in g.FIRMA_POLA] + [(g.POLE_ID, True), (g.POLE_STRONA_CEN, True)]
    for r, (p, wym) in enumerate(pola, start=2):
        wiersz(ws, r, [p, przyklad.get(p)], fill=WYM if wym else OPC)
        ws.cell(row=r, column=2).number_format = "@"
    ws.cell(row=len(pola) + 3, column=1, value="Żółte = wymagane, szare = opcjonalne. KRS albo CEiDG: wystarczy jedno.").font = Font(name=F, italic=True)
    ws.cell(row=pola.index((g.POLE_ID, True)) + 2, column=1).comment = Comment(
        "Krótka nazwa bez spacji i polskich znaków, np. nazwa-firmy. Trafia do adresów plików i identyfikatorów "
        "na dane.gov.pl. Po pierwszej publikacji NIE ZMIENIAĆ.", "Automat")

    # ---------------- Inwestycje
    ws = wb.create_sheet("Inwestycje")
    kol = [k for k, *_ in g.INWESTYCJE_KOLUMNY]
    naglowek(ws, kol, [24, 16, 16, 18, 18, 22, 14, 13, 44])
    wiersz(ws, 2, ["Osiedle Zielone", "wielkopolskie", "poznański", "Komorniki", "Komorniki", "ul. Polna",
                   "12", "62-052", "https://przyklad-deweloper.pl/osiedle-zielone/prospekt"])
    for c in range(1, 10):
        ws.cell(row=2, column=c).number_format = "@"

    # ---------------- Oferta
    ws = wb.create_sheet("Oferta")
    naglowek(ws, g.OFERTA_KOLUMNY, [18, 22, 11, 12, 15, 15, 20, 12, 12, 22, 13, 13, 22, 13, 22, 13, 18])
    przyk = [
        ["Osiedle Zielone", "Lokal mieszkalny", "A1", 12500, 687500, None, "Miejsce postojowe", "MP1", 35000,
         "Komórka lokatorska", "KL1", 12000, None, None, None, None, None],
        ["Osiedle Zielone", "Lokal mieszkalny", "A2", 12900, 812700, None, None, None, None, None, None, None,
         None, None, None, None, None],
        ["Osiedle Zielone", None, None, None, None, None, "Miejsce postojowe", "MP7", 39000, None, None, None,
         None, None, None, None, None],
    ]
    for r, w in enumerate(przyk, start=2):
        wiersz(ws, r, w)
        ws.cell(row=r, column=3).number_format = "@"
        for c in (4, 5, 6, 9, 12, 14, 16):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
    dv = DataValidation(type="list", formula1='"Lokal mieszkalny,Dom jednorodzinny"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B2:B2000")
    dvi = DataValidation(type="list", formula1="=Inwestycje!$A$2:$A$100", allow_blank=False)
    ws.add_data_validation(dvi)
    dvi.add("A2:A2000")
    for c in range(1, 18):
        ws.cell(row=1, column=c).alignment = Alignment(wrap_text=True, vertical="center")

    wb.active = 3
    wb.save(cel)
    print("Zapisano", cel)


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "Cennik_szablon.xlsx")
