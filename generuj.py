#!/usr/bin/env python3
"""
Generator codziennego raportu cen mieszkań dla dane.gov.pl (art. 19b ustawy deweloperskiej).

Co robi przy każdym uruchomieniu:
  1. Pobiera cennik (Google Sheets jako .xlsx albo lokalny plik Excel) z 3 zakładkami:
     "Firma", "Inwestycje", "Oferta".
  2. Sprawdza dane i przerywa z czytelnym komunikatem, jeśli czegoś brakuje.
  3. Sam pilnuje dat zmian cen (plik stan/ceny.json): gdy cena się zmienia, data = dzień zmiany.
  4. Zapisuje dzisiejszy plik CSV w formacie „Wzorcowy zakres danych” Ministerstwa Cyfryzacji (58 kolumn).
  5. Przebudowuje plik XML dla Importera XML dane.gov.pl (wszystkie dni, zbiór danych na każdy rok)
     oraz plik .md5 z sumą kontrolną (małe litery, bez nazwy pliku i bez nowej linii).

Wymaga: Python 3.9+, openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import sys
import time
import unicodedata
import urllib.request
from pathlib import Path
from xml.sax.saxutils import escape

try:
    from zoneinfo import ZoneInfo
except ImportError:  # pragma: no cover
    ZoneInfo = None

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
NAGLOWKI = json.loads((Path(__file__).parent / "naglowki_ministerstwo.json").read_text(encoding="utf-8"))
assert len(NAGLOWKI) == 58, "Plik naglowki_ministerstwo.json musi mieć 58 kolumn"

X = "X"  # znak umowny: wypełnienie pozycji niemożliwe lub niecelowe
NS = "urn:otwarte-dane:harvester:1.13"

# ----------------------------------------------------------------------------------------------
# Opis arkusza wejściowego
# ----------------------------------------------------------------------------------------------
# Zakładka "Firma": kolumna A = nazwa pola, kolumna B = wartość.  (pole, indeks kolumny w CSV, typ, wymagane)
FIRMA_POLA = [
    ("Nazwa dewelopera", 0, "tekst", True),
    ("Forma prawna", 1, "tekst", True),
    ("Nr KRS", 2, "krs", False),
    ("Nr wpisu do CEiDG", 3, "tekst", False),
    ("Nr NIP", 4, "nip", True),
    ("Nr REGON", 5, "regon", True),
    ("Nr telefonu", 6, "telefon", True),
    ("Adres e-mail", 7, "tekst", True),
    ("Nr faxu", 8, "telefon", False),
    ("Adres strony internetowej", 9, "tekst", True),
    ("Siedziba - województwo", 10, "tekst", True),
    ("Siedziba - powiat", 11, "tekst", True),
    ("Siedziba - gmina", 12, "tekst", True),
    ("Siedziba - miejscowość", 13, "tekst", True),
    ("Siedziba - ulica", 14, "tekst", False),
    ("Siedziba - nr nieruchomości", 15, "tekst", True),
    ("Siedziba - nr lokalu", 16, "tekst", False),
    ("Siedziba - kod pocztowy", 17, "kod", True),
    ("Biuro sprzedaży - województwo", 18, "tekst", False),
    ("Biuro sprzedaży - powiat", 19, "tekst", False),
    ("Biuro sprzedaży - gmina", 20, "tekst", False),
    ("Biuro sprzedaży - miejscowość", 21, "tekst", False),
    ("Biuro sprzedaży - ulica", 22, "tekst", False),
    ("Biuro sprzedaży - nr nieruchomości", 23, "tekst", False),
    ("Biuro sprzedaży - nr lokalu", 24, "tekst", False),
    ("Biuro sprzedaży - kod pocztowy", 25, "kod", False),
    ("Dodatkowe lokalizacje sprzedaży", 26, "tekst", False),
    ("Sposób kontaktu nabywcy z deweloperem", 27, "tekst", True),
]
# Pola techniczne (nie trafiają do CSV)
POLE_ID = "Identyfikator techniczny (bez spacji, NIE ZMIENIAĆ po starcie)"
POLE_STRONA_CEN = "Adres strony z cenami (URL)"

INWESTYCJE_KOLUMNY = [
    # (nagłówek w arkuszu, indeks w CSV albo None, typ, wymagane)
    ("Inwestycja", None, "tekst", True),
    ("Województwo", 28, "tekst", True),
    ("Powiat", 29, "tekst", True),
    ("Gmina", 30, "tekst", True),
    ("Miejscowość", 31, "tekst", True),
    ("Ulica", 32, "tekst", False),
    ("Nr nieruchomości / działki", 33, "tekst", False),
    ("Kod pocztowy", 34, "kod", True),
    ("Adres prospektu informacyjnego (URL)", 57, "tekst", True),
]

OFERTA_KOLUMNY = [
    "Inwestycja",
    "Rodzaj nieruchomości (Lokal mieszkalny / Dom jednorodzinny)",
    "Nr lokalu lub domu",
    "Cena za m² [zł]",
    "Cena lokalu: m² × powierzchnia [zł]",
    "Cena łączna z innymi składnikami [zł]",
    "Rodzaj części nieruchomości (np. Miejsce postojowe)",
    "Oznaczenie części",
    "Cena części [zł]",
    "Rodzaj pomieszczenia przynależnego (np. Komórka lokatorska)",
    "Oznaczenie pomieszczenia",
    "Cena pomieszczenia [zł]",
    "Prawa niezbędne - opis",
    "Prawa niezbędne - wartość [zł]",
    "Inne świadczenia pieniężne - opis",
    "Inne świadczenia pieniężne - wartość [zł]",
    "Ceny obowiązują od (opcjonalnie, tylko przy pierwszym wpisie)",
]

# pole ceny w ofercie -> (indeks ceny w CSV, indeks daty w CSV)
CENY = {
    "Cena za m² [zł]": (37, 38),
    "Cena lokalu: m² × powierzchnia [zł]": (39, 40),
    "Cena łączna z innymi składnikami [zł]": (41, 42),
    "Cena części [zł]": (45, 46),
    "Cena pomieszczenia [zł]": (49, 50),
    "Prawa niezbędne - wartość [zł]": (52, 53),
    "Inne świadczenia pieniężne - wartość [zł]": (55, 56),
}
TEKSTY_OFERTY = {
    "Rodzaj nieruchomości (Lokal mieszkalny / Dom jednorodzinny)": 35,
    "Nr lokalu lub domu": 36,
    "Rodzaj części nieruchomości (np. Miejsce postojowe)": 43,
    "Oznaczenie części": 44,
    "Rodzaj pomieszczenia przynależnego (np. Komórka lokatorska)": 47,
    "Oznaczenie pomieszczenia": 48,
    "Prawa niezbędne - opis": 51,
    "Inne świadczenia pieniężne - opis": 54,
}
RODZAJE = {"lokal mieszkalny": "Lokal mieszkalny", "dom jednorodzinny": "Dom jednorodzinny"}


class BladDanych(Exception):
    pass


# ----------------------------------------------------------------------------------------------
# Pomocnicze
# ----------------------------------------------------------------------------------------------
def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def pusty(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() in ("", "-", "—"))


def tekst(v) -> str:
    if pusty(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return re.sub(r"\s+", " ", str(v)).strip()


def cyfry(v) -> str:
    return re.sub(r"\D", "", tekst(v))


def liczba(v, gdzie: str) -> str:
    """Zwraca kwotę jako tekst z kropką dziesiętną, max 2 miejsca, bez separatorów tysięcy."""
    if pusty(v) or (isinstance(v, str) and v.strip().upper() == X):
        return ""
    if isinstance(v, (int, float)):
        x = float(v)
    else:
        s = str(v).replace("zł", "").replace("PLN", "").replace("\xa0", "").replace(" ", "")
        if s.count(",") == 1 and s.count(".") == 0:
            s = s.replace(",", ".")
        elif s.count(",") >= 1 and s.count(".") == 1:  # 1,234.56
            s = s.replace(",", "")
        try:
            x = float(s)
        except ValueError:
            raise BladDanych(f"{gdzie}: „{v}” nie jest liczbą")
    if x < 0:
        raise BladDanych(f"{gdzie}: cena nie może być ujemna ({v})")
    x = round(x, 2)
    return str(int(x)) if x.is_integer() else f"{x:.2f}"


def data_iso(v, gdzie: str) -> str:
    if pusty(v):
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    for f in ("%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y", "%Y.%m.%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s, f).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise BladDanych(f"{gdzie}: nie rozumiem daty „{v}” (użyj RRRR-MM-DD)")


def slug(s: str) -> str:
    s = unicodedata.normalize("NFKD", s.replace("ł", "l").replace("Ł", "L"))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def dzis_warszawa() -> dt.date:
    if ZoneInfo:
        return dt.datetime.now(ZoneInfo("Europe/Warsaw")).date()
    return dt.date.today()


def sprawdz_format(wartosc: str, typ: str, gdzie: str, bledy: list) -> str:
    if not wartosc:
        return wartosc
    if typ == "krs":
        d = cyfry(wartosc).zfill(10)
        if len(d) != 10:
            bledy.append(f"{gdzie}: KRS powinien mieć 10 cyfr")
        return d
    if typ == "nip":
        d = cyfry(wartosc)
        if len(d) != 10:
            bledy.append(f"{gdzie}: NIP powinien mieć 10 cyfr")
        return d
    if typ == "regon":
        d = cyfry(wartosc)
        if len(d) not in (9, 14):
            bledy.append(f"{gdzie}: REGON powinien mieć 9 lub 14 cyfr")
        return d
    if typ == "telefon":
        plus = wartosc.strip().startswith("+")
        d = cyfry(wartosc)
        return ("+" if plus else "") + d
    if typ == "kod":
        d = cyfry(wartosc)
        if len(d) != 5:
            bledy.append(f"{gdzie}: kod pocztowy powinien mieć format 00-000")
            return wartosc
        return f"{d[:2]}-{d[2:]}"
    return wartosc


# ----------------------------------------------------------------------------------------------
# Wczytanie arkusza
# ----------------------------------------------------------------------------------------------
def pobierz_zrodlo(zrodlo: str, cel: Path) -> Path:
    if re.match(r"^https?://", zrodlo):
        ostatni = None
        for proba in range(1, 4):
            try:
                req = urllib.request.Request(zrodlo, headers={"User-Agent": "jawnosc-cen-automat"})
                with urllib.request.urlopen(req, timeout=60) as r:
                    dane = r.read()
                if not dane.startswith(b"PK"):
                    raise BladDanych(
                        "Pobrany plik nie jest arkuszem Excel. Sprawdź, czy arkusz Google jest udostępniony "
                        "jako „Każda osoba mająca link: przeglądający”."
                    )
                cel.write_bytes(dane)
                return cel
            except BladDanych:
                raise
            except Exception as e:  # sieć
                ostatni = e
                time.sleep(10 * proba)
        raise ConnectionError(f"Nie udało się pobrać arkusza: {ostatni}")
    p = Path(zrodlo)
    if not p.exists():
        raise BladDanych(f"Nie znaleziono pliku z cennikiem: {p}")
    return p


def arkusz(wb, nazwa):
    for ws in wb.worksheets:
        if norm(ws.title) == norm(nazwa):
            return ws
    raise BladDanych(f"W pliku brakuje zakładki „{nazwa}”")


def wiersze_tabeli(ws, oczekiwane: list[str], nazwa: str):
    """Zwraca listę (nr_wiersza, dict) – kolumny dopasowane po nagłówkach z wiersza 1."""
    naglowki = [norm(c.value) for c in ws[1]]
    mapa = {}
    for kol in oczekiwane:
        if norm(kol) in naglowki:
            mapa[kol] = naglowki.index(norm(kol))
        elif norm(kol).startswith("ceny obowiązują od"):
            mapa[kol] = None  # kolumna opcjonalna
        else:
            raise BladDanych(f"Zakładka „{nazwa}”: brak kolumny „{kol}” w pierwszym wierszu")
    wynik = []
    for nr, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        d = {k: (row[i] if i is not None and i < len(row) else None) for k, i in mapa.items()}
        if all(pusty(v) for v in d.values()):
            continue
        wynik.append((nr, d))
    return wynik


def wczytaj(sciezka: Path):
    wb = openpyxl.load_workbook(sciezka, data_only=True)
    bledy: list[str] = []

    # --- Firma
    ws = arkusz(wb, "Firma")
    kv = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if row and not pusty(row[0]):
            kv[norm(row[0])] = row[1] if len(row) > 1 else None
    firma = [""] * 58
    for pole, idx, typ, wym in FIRMA_POLA:
        w = tekst(kv.get(norm(pole)))
        if wym and not w:
            bledy.append(f"Firma: uzupełnij pole „{pole}”")
        firma[idx] = sprawdz_format(w, typ, f"Firma / {pole}", bledy)
    if not firma[2] and not firma[3]:
        bledy.append("Firma: podaj Nr KRS (spółka) albo Nr wpisu do CEiDG (jednoosobowa działalność)")
    ident = slug(tekst(kv.get(norm(POLE_ID))))
    if not ident:
        bledy.append(f"Firma: uzupełnij pole „{POLE_ID}” (np. nazwa-firmy)")
    strona_cen = tekst(kv.get(norm(POLE_STRONA_CEN))) or firma[9]
    meta = {"id": ident[:24], "nazwa": firma[0], "strona_cen": strona_cen}

    # --- Inwestycje
    inwestycje = {}
    for nr, d in wiersze_tabeli(arkusz(wb, "Inwestycje"), [k for k, *_ in INWESTYCJE_KOLUMNY], "Inwestycje"):
        nazwa = tekst(d["Inwestycja"])
        kol = [""] * 58
        for k, idx, typ, wym in INWESTYCJE_KOLUMNY:
            w = tekst(d[k])
            if wym and not w:
                bledy.append(f"Inwestycje, wiersz {nr}: uzupełnij „{k}”")
            if idx is not None:
                kol[idx] = sprawdz_format(w, typ, f"Inwestycje, wiersz {nr}", bledy)
        if norm(nazwa) in inwestycje:
            bledy.append(f"Inwestycje, wiersz {nr}: inwestycja „{nazwa}” występuje dwa razy")
        inwestycje[norm(nazwa)] = kol
    if not inwestycje:
        bledy.append("Inwestycje: dodaj co najmniej jedną inwestycję")

    # --- Oferta
    oferta = []
    klucze = set()
    for nr, d in wiersze_tabeli(arkusz(wb, "Oferta"), OFERTA_KOLUMNY, "Oferta"):
        g = f"Oferta, wiersz {nr}"
        try:
            inw = tekst(d["Inwestycja"])
            if norm(inw) not in inwestycje:
                bledy.append(f"{g}: inwestycja „{inw}” nie istnieje w zakładce Inwestycje")
                continue
            teksty = {k: tekst(d[k]) for k in TEKSTY_OFERTY}
            ceny = {k: liczba(d[k], f"{g} / {k}") for k in CENY}
            od = data_iso(d.get(OFERTA_KOLUMNY[-1]), f"{g} / data obowiązywania")
        except BladDanych as e:
            bledy.append(str(e))
            continue
        rodzaj_k = OFERTA_KOLUMNY[1]
        if teksty[rodzaj_k]:
            r = RODZAJE.get(norm(teksty[rodzaj_k]))
            if not r:
                bledy.append(f"{g}: rodzaj nieruchomości musi być „Lokal mieszkalny” albo „Dom jednorodzinny”")
                continue
            teksty[rodzaj_k] = r
            if not teksty["Nr lokalu lub domu"]:
                bledy.append(f"{g}: brak numeru lokalu/domu")
            if not ceny["Cena za m² [zł]"]:
                bledy.append(f"{g}: brak ceny za m²")
            if not (ceny["Cena lokalu: m² × powierzchnia [zł]"] or ceny["Cena łączna z innymi składnikami [zł]"]):
                bledy.append(f"{g}: podaj cenę lokalu (m² × powierzchnia) lub cenę łączną")
        elif not (teksty["Oznaczenie części"] or teksty["Oznaczenie pomieszczenia"]):
            bledy.append(f"{g}: wiersz bez lokalu musi mieć oznaczenie części (np. MP1) lub pomieszczenia (np. KL1)")
        for opis, cena in [("Rodzaj części nieruchomości (np. Miejsce postojowe)", "Cena części [zł]"),
                           ("Rodzaj pomieszczenia przynależnego (np. Komórka lokatorska)", "Cena pomieszczenia [zł]"),
                           ("Prawa niezbędne - opis", "Prawa niezbędne - wartość [zł]"),
                           ("Inne świadczenia pieniężne - opis", "Inne świadczenia pieniężne - wartość [zł]")]:
            if ceny[cena] and not teksty[opis]:
                bledy.append(f"{g}: podano „{cena}”, ale brak opisu w „{opis}”")
        klucz = "|".join(norm(x) for x in (inw, teksty["Nr lokalu lub domu"], teksty["Oznaczenie części"],
                                           teksty["Oznaczenie pomieszczenia"]))
        if klucz in klucze:
            bledy.append(f"{g}: ta sama pozycja (inwestycja + numer + oznaczenia) występuje drugi raz")
        klucze.add(klucz)
        oferta.append({"nr": nr, "inw": norm(inw), "klucz": klucz, "teksty": teksty, "ceny": ceny, "od": od})
    if not oferta and os.environ.get("POZWOL_PUSTA_OFERTE") != "1":
        bledy.append("Oferta: brak pozycji. Jeśli naprawdę nic nie jest w sprzedaży, ustaw POZWOL_PUSTA_OFERTE=1")

    if bledy:
        raise BladDanych("Znaleziono błędy w cenniku:\n  - " + "\n  - ".join(bledy))
    return meta, firma, inwestycje, oferta


# ----------------------------------------------------------------------------------------------
# Budowa CSV
# ----------------------------------------------------------------------------------------------
def zbuduj_wiersze(firma, inwestycje, oferta, stan: dict, dzien: str):
    wiersze = []
    nowy_stan = {}
    for poz in oferta:
        w = list(firma)
        inw = inwestycje[poz["inw"]]
        for i in list(range(28, 35)) + [57]:
            w[i] = inw[i]
        for k, idx in TEKSTY_OFERTY.items():
            w[idx] = poz["teksty"][k]
        stary = stan.get(poz["klucz"], {})
        nowy = {}
        for k, (ic, idd) in CENY.items():
            cena = poz["ceny"][k]
            if not cena:
                continue
            poprzednia = stary.get(k)
            if poprzednia and poprzednia[0] == cena:
                od = poprzednia[1]
            elif not poprzednia and poz["od"]:
                od = poz["od"]
            else:
                od = dzien
            w[ic], w[idd] = cena, od
            nowy[k] = [cena, od]
        nowy_stan[poz["klucz"]] = nowy
        wiersze.append(["X" if v in ("", None) else v for v in w])
    return wiersze, nowy_stan


def zapisz_csv(sciezka: Path, wiersze):
    sciezka.parent.mkdir(parents=True, exist_ok=True)
    with open(sciezka, "w", encoding="utf-8", newline="") as f:
        wr = csv.writer(f, delimiter=",", quotechar='"', quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
        wr.writerow(NAGLOWKI)
        wr.writerows(wiersze)


# ----------------------------------------------------------------------------------------------
# Budowa XML + MD5
# ----------------------------------------------------------------------------------------------
USTAWA = ("art. 19b ust. 1 ustawy z dnia 20 maja 2021 r. o ochronie praw nabywcy lokalu mieszkalnego lub domu "
          "jednorodzinnego oraz Deweloperskim Funduszu Gwarancyjnym (Dz. U. z 2024 r. poz. 695)")


def przytnij(s: str, n: int) -> str:
    return s if len(s) <= n else s[: n - 1] + "…"


def nazwa_csv(ident: str, dzien: str) -> str:
    return f"ceny-ofertowe-mieszkan-dewelopera-{ident}-{dzien}.csv"


def zbuduj_xml(meta, katalog: Path, base_url: str) -> str:
    wzor = re.compile(rf"^ceny-ofertowe-mieszkan-dewelopera-{re.escape(meta['id'])}-(\d{{4}}-\d{{2}}-\d{{2}})\.csv$")
    pliki = {}
    for p in sorted((katalog / "csv").glob("*/*.csv")):
        m = wzor.match(p.name)
        if m:
            pliki[m.group(1)] = p
    lata = sorted({d[:4] for d in pliki})
    n = escape(meta["nazwa"])
    L = ['<?xml version="1.0" encoding="UTF-8"?>',
         f'<ns2:datasets xmlns:ns2="{NS}" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">']
    for rok in lata:
        L += [
            '  <dataset status="published">',
            f"    <extIdent>{meta['id']}-ceny-{rok}</extIdent>",
            "    <title>",
            f"      <polish>{przytnij(f'Ceny ofertowe mieszkań dewelopera {n} w {rok} r.', 300)}</polish>",
            f"      <english>{przytnij(f'Offer prices of apartments of developer {n} in {rok}', 300)}</english>",
            "    </title>",
            "    <description>",
            f"      <polish>{przytnij(f'Zbiór danych zawiera informacje o cenach ofertowych mieszkań dewelopera {n} udostępniane zgodnie z {USTAWA}.', 1000)}</polish>",
            f"      <english>{przytnij(f'The dataset contains information on offer prices of apartments of the developer {n} made available in accordance with {USTAWA}.', 1000)}</english>",
            "    </description>",
            f"    <url>{escape(meta['strona_cen'])}</url>",
            "    <updateFrequency>daily</updateFrequency>",
            "    <hasDynamicData>false</hasDynamicData>",
            "    <hasHighValueData>true</hasHighValueData>",
            "    <hasHighValueDataFromEuropeanCommissionList>false</hasHighValueDataFromEuropeanCommissionList>",
            "    <hasResearchData>false</hasResearchData>",
            "    <categories>",
            "      <category>ECON</category>",
            "    </categories>",
            "    <resources>",
        ]
        for dzien in sorted(d for d in pliki if d.startswith(rok)):
            url = f"{base_url.rstrip('/')}/csv/{rok}/{nazwa_csv(meta['id'], dzien)}"
            L += [
                '      <resource status="published">',
                f"        <extIdent>{meta['id']}-{dzien.replace('-', '')}</extIdent>",
                f"        <url>{escape(url)}</url>",
                "        <title>",
                f"          <polish>{przytnij(f'Ceny ofertowe mieszkań dewelopera {n} {dzien}', 300)}</polish>",
                f"          <english>{przytnij(f'Offer prices of apartments of developer {n} {dzien}', 300)}</english>",
                "        </title>",
                "        <description>",
                f"          <polish>{przytnij(f'Dane dotyczące cen ofertowych mieszkań dewelopera {n} udostępnione {dzien} zgodnie z {USTAWA}.', 1000)}</polish>",
                f"          <english>{przytnij(f'Data on offer prices of apartments of the developer {n} made available {dzien} in accordance with {USTAWA}.', 1000)}</english>",
                "        </description>",
                "        <availability>local</availability>",
                f"        <dataDate>{dzien}</dataDate>",
                "        <specialSigns>",
                "          <specialSign>X</specialSign>",
                "        </specialSigns>",
                "        <hasDynamicData>false</hasDynamicData>",
                "        <hasHighValueData>true</hasHighValueData>",
                "        <hasHighValueDataFromEuropeanCommissionList>false</hasHighValueDataFromEuropeanCommissionList>",
                "        <hasResearchData>false</hasResearchData>",
                "        <containsProtectedData>false</containsProtectedData>",
                "      </resource>",
            ]
        L += [
            "    </resources>",
            "    <tags>",
            '      <tag lang="pl">Deweloper</tag>',
            "    </tags>",
            "  </dataset>",
        ]
    L.append("</ns2:datasets>")
    return "\n".join(L) + "\n"


def zapisz_xml_md5(meta, katalog: Path, base_url: str):
    xml = zbuduj_xml(meta, katalog, base_url).encode("utf-8")
    (katalog / f"{meta['id']}.xml").write_bytes(xml)
    (katalog / f"{meta['id']}.md5").write_text(hashlib.md5(xml).hexdigest(), encoding="ascii")
    return xml


def zapisz_index(meta, katalog: Path, base_url: str, dzien: str, liczba_pozycji: int, awaryjnie: bool):
    xml_url = f"{base_url.rstrip('/')}/{meta['id']}.xml"
    uwaga = "<p style='color:#b00'><b>Uwaga:</b> dziś opublikowano kopię ostatniego cennika (arkusz był niedostępny).</p>" if awaryjnie else ""
    html = f"""<!doctype html><html lang="pl"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Raport cen – {escape(meta['nazwa'])}</title>
<body style="font-family:system-ui,sans-serif;max-width:720px;margin:40px auto;padding:0 16px;line-height:1.5">
<h1>Raport cen ofertowych – {escape(meta['nazwa'])}</h1>
<p>Ostatnia publikacja: <b>{dzien}</b> · pozycji w ofercie: <b>{liczba_pozycji}</b></p>{uwaga}
<p>Plik dla dane.gov.pl (Importer XML): <a href="{escape(xml_url)}">{escape(xml_url)}</a></p>
<p>Dzisiejszy cennik CSV: <a href="csv/{dzien[:4]}/{nazwa_csv(meta['id'], dzien)}">{nazwa_csv(meta['id'], dzien)}</a></p>
</body></html>
"""
    (katalog / "index.html").write_text(html, encoding="utf-8")
    (katalog / ".nojekyll").write_text("", encoding="utf-8")


# ----------------------------------------------------------------------------------------------
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--zrodlo", help="ścieżka do .xlsx albo URL; domyślnie arkusz Google z GOOGLE_SHEET_ID lub cennik.xlsx")
    ap.add_argument("--data", help="data raportu RRRR-MM-DD (domyślnie dziś wg czasu warszawskiego)")
    ap.add_argument("--wyjscie", default=str(ROOT / "docs"), help="katalog publikowany na stronie (GitHub Pages)")
    ap.add_argument("--stan", default=str(ROOT / "stan"), help="katalog z historią cen")
    ap.add_argument("--base-url", default=os.environ.get("PAGES_URL", ""), help="publiczny adres katalogu docs")
    ap.add_argument("--tylko-sprawdz", action="store_true", help="tylko sprawdź arkusz, nic nie zapisuj")
    a = ap.parse_args(argv)

    dzien = a.data or dzis_warszawa().isoformat()
    wyj, stan_dir = Path(a.wyjscie), Path(a.stan)
    stan_dir.mkdir(parents=True, exist_ok=True)
    wyj.mkdir(parents=True, exist_ok=True)
    plik_stanu = stan_dir / "ceny.json"
    plik_meta = stan_dir / "meta.json"

    zrodlo = a.zrodlo
    if not zrodlo and os.environ.get("GOOGLE_SHEET_ID"):
        zrodlo = f"https://docs.google.com/spreadsheets/d/{os.environ['GOOGLE_SHEET_ID'].strip()}/export?format=xlsx"
    if not zrodlo:
        zrodlo = str(ROOT / "cennik.xlsx")

    if not a.base_url and not a.tylko_sprawdz:
        print("BŁĄD: brak publicznego adresu (PAGES_URL / --base-url)", file=sys.stderr)
        return 2

    awaryjnie = False
    try:
        plik = pobierz_zrodlo(zrodlo, stan_dir / "_pobrany_cennik.xlsx")
        meta, firma, inwestycje, oferta = wczytaj(plik)
    except BladDanych as e:
        print(f"::error::{e}" if os.environ.get("GITHUB_ACTIONS") else f"BŁĄD: {e}", file=sys.stderr)
        return 2
    except ConnectionError as e:
        # Arkusz chwilowo niedostępny: publikujemy kopię ostatniego cennika, żeby nie przerwać ciągłości.
        if a.tylko_sprawdz or not plik_meta.exists():
            print(f"BŁĄD: {e}", file=sys.stderr)
            return 2
        meta = json.loads(plik_meta.read_text(encoding="utf-8"))
        ostatnie = sorted((wyj / "csv").glob(f"*/ceny-ofertowe-mieszkan-dewelopera-{meta['id']}-*.csv"))
        if not ostatnie:
            print(f"BŁĄD: {e}", file=sys.stderr)
            return 2
        cel = wyj / "csv" / dzien[:4] / nazwa_csv(meta["id"], dzien)
        if ostatnie[-1] != cel:
            cel.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(ostatnie[-1], cel)
        awaryjnie, oferta = True, []
        print(f"::warning::{e}. Opublikowano kopię cennika z pliku {ostatnie[-1].name}.", file=sys.stderr)
    finally:
        tmp = stan_dir / "_pobrany_cennik.xlsx"
        if tmp.exists():
            tmp.unlink()

    if a.tylko_sprawdz:
        print(f"OK – arkusz poprawny: {len(inwestycje)} inwestycji, {len(oferta)} pozycji w ofercie.")
        return 0

    if not awaryjnie:
        if plik_meta.exists():
            stare = json.loads(plik_meta.read_text(encoding="utf-8"))
            if stare.get("id") and stare["id"] != meta["id"]:
                print(f"BŁĄD: zmieniono identyfikator techniczny z „{stare['id']}” na „{meta['id']}”. "
                      "To utworzyłoby nowy zbiór na dane.gov.pl. Przywróć poprzednią wartość w zakładce Firma.",
                      file=sys.stderr)
                return 2
        stan = json.loads(plik_stanu.read_text(encoding="utf-8")) if plik_stanu.exists() else {}
        wiersze, nowy_stan = zbuduj_wiersze(firma, inwestycje, oferta, stan, dzien)
        # pozycje, które zniknęły z oferty, zostają w stanie (na wypadek powrotu do sprzedaży)
        stan.update(nowy_stan)
        zapisz_csv(wyj / "csv" / dzien[:4] / nazwa_csv(meta["id"], dzien), wiersze)
        plik_stanu.write_text(json.dumps(stan, ensure_ascii=False, indent=1, sort_keys=True), encoding="utf-8")
        plik_meta.write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")

    zapisz_xml_md5(meta, wyj, a.base_url)
    zapisz_index(meta, wyj, a.base_url, dzien, len(oferta), awaryjnie)
    print(f"OK – {dzien}: opublikowano {nazwa_csv(meta['id'], dzien)} ({'kopia awaryjna' if awaryjnie else f'{len(oferta)} pozycji'}).")
    print(f"     Adres dla dane.gov.pl: {a.base_url.rstrip('/')}/{meta['id']}.xml")
    return 3 if awaryjnie else 0


if __name__ == "__main__":
    sys.exit(main())
